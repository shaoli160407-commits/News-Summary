from __future__ import annotations

import argparse
import json
import os
import smtplib
import ssl
import sys
import urllib.error
import urllib.request
from datetime import date, datetime
from email.message import EmailMessage
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from dotenv import load_dotenv


if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")


REQUIRED_ITEM_FIELDS = ("title", "source", "date", "url", "summary", "impact")
DEFAULT_DEEPSEEK_BASE_URL = "https://api.deepseek.com/anthropic"
DEFAULT_DEEPSEEK_MODEL = "deepseek-v4-pro"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build and optionally email a global trade daily DOCX report.")
    parser.add_argument("payload", nargs="?", type=Path, help="Optional UTF-8 JSON payload path.")
    parser.add_argument("--deepseek-research", action="store_true", help="Ask DeepSeek API to research and generate the payload JSON.")
    parser.add_argument("--report-date", default=date.today().isoformat(), help="Report date in YYYY-MM-DD format.")
    parser.add_argument("--days", type=int, default=5, help="Recent-news window in days.")
    parser.add_argument("--output-dir", type=Path, default=Path("output"), help="Directory for the generated DOCX.")
    parser.add_argument("--send-email", action="store_true", help="Send the DOCX to GMAIL_ADDRESS using Gmail SMTP.")
    parser.add_argument("--env-file", type=Path, default=Path(".env"), help="Path to .env containing API and Gmail credentials.")
    parser.add_argument("--save-payload", type=Path, help="Optional path to save the DeepSeek-generated payload JSON.")
    parser.add_argument("--deepseek-model", default=DEFAULT_DEEPSEEK_MODEL, help="DeepSeek model name.")
    parser.add_argument("--deepseek-base-url", default=DEFAULT_DEEPSEEK_BASE_URL, help="DeepSeek Anthropic-compatible base URL.")
    parser.add_argument("--max-tokens", type=int, default=12000, help="Max output tokens for DeepSeek.")
    parser.add_argument("--max-searches", type=int, default=12, help="Max web-search uses for DeepSeek server tool.")
    return parser.parse_args()


def read_text_response(response: dict[str, Any]) -> str:
    blocks = response.get("content", [])
    texts: list[str] = []
    for block in blocks:
        if isinstance(block, dict) and block.get("type") == "text":
            texts.append(str(block.get("text", "")))
    if texts:
        return "\n".join(texts).strip()
    return json.dumps(response, ensure_ascii=False)


def extract_json_object(text: str) -> dict[str, Any]:
    stripped = text.strip()
    if stripped.startswith("```"):
        lines = stripped.splitlines()
        if lines and lines[0].startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].startswith("```"):
            lines = lines[:-1]
        stripped = "\n".join(lines).strip()
    try:
        return json.loads(stripped)
    except json.JSONDecodeError:
        start = stripped.find("{")
        end = stripped.rfind("}")
        if start < 0 or end <= start:
            raise ValueError("DeepSeek 返回内容中没有可解析的 JSON 对象")
        return json.loads(stripped[start:end + 1])


def deepseek_key(env_file: Path) -> str:
    load_dotenv(env_file)
    key = os.getenv("DEEPSEEK_API_KEY", "").strip()
    if not key:
        raise RuntimeError("缺少 DeepSeek 配置：DEEPSEEK_API_KEY")
    return key


def deepseek_prompt(report_date: str, days: int) -> str:
    return f"""
请调用可用的网页搜索能力，检索并总结截至 {report_date} 的最近 {days} 天全球外贸相关新闻与市场动态。

搜索关键词必须覆盖：
- 中文：全球外贸 最新动态、进出口政策、关税 贸易、供应链 外贸、汇率 出口
- 英文：global trade news、import export policy、tariffs trade war、supply chain disruption、trade sanctions

筛选标准：
- 只保留对外贸市场有实质影响的事件。
- 优先：关税政策调整、贸易制裁、进出口数据、汇率波动、大宗商品价格、供应链中断或恢复、重大贸易协定。
- 至少 10 条新闻，每条必须有真实来源、日期、URL。
- 不要编造来源或 URL；如果某条没有可靠来源，不要使用。

输出要求：
- 只输出一个合法 JSON 对象，不要 Markdown，不要解释。
- JSON 必须匹配以下结构：
{{
  "report_date": "{report_date}",
  "title": "全球外贸资讯日报",
  "executive_summary": "200字以内中文摘要",
  "trend": "100字以内中文趋势研判",
  "categories": [
    {{
      "name": "政策法规",
      "items": [
        {{
          "title": "新闻标题",
          "source": "来源名称",
          "date": "YYYY-MM-DD",
          "url": "https://...",
          "summary": "3-5句话核心内容",
          "impact": "对外贸的潜在影响"
        }}
      ]
    }},
    {{"name": "市场行情", "items": []}},
    {{"name": "供应链与物流", "items": []}},
    {{"name": "汇率与金融", "items": []}}
  ]
}}

分类要求：
- 每类尽量 3-5 条；如果某类当天有效新闻不足，可以少于 3 条，但总数至少 10 条。
- 日期必须在最近 {days} 天窗口内，除非是仍在持续影响市场的事件并在窗口内有新进展。
""".strip()


def call_deepseek_research(args: argparse.Namespace) -> dict[str, Any]:
    key = deepseek_key(args.env_file)
    url = args.deepseek_base_url.rstrip("/") + "/v1/messages"
    body = {
        "model": args.deepseek_model,
        "max_tokens": args.max_tokens,
        "temperature": 0.2,
        "system": "你是严谨的国际贸易资讯研究员。必须联网检索，必须保留来源链接，必须输出合法 JSON。",
        "messages": [
            {
                "role": "user",
                "content": [{"type": "text", "text": deepseek_prompt(args.report_date, args.days)}],
            }
        ],
        "tools": [
            {
                "type": "web_search_20250305",
                "name": "web_search",
                "max_uses": args.max_searches,
            }
        ],
        "tool_choice": {"type": "auto"},
    }
    request = urllib.request.Request(
        url,
        data=json.dumps(body, ensure_ascii=False).encode("utf-8"),
        headers={
            "content-type": "application/json",
            "x-api-key": key,
            "anthropic-version": "2023-06-01",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=180) as response:
            raw = response.read().decode("utf-8")
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"DeepSeek API 请求失败：HTTP {exc.code} {detail}") from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(f"DeepSeek API 网络请求失败：{exc}") from exc

    response_data = json.loads(raw)
    text = read_text_response(response_data)
    payload = extract_json_object(text)
    validate_payload(payload)
    return payload


def load_payload(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(f"找不到日报数据文件：{path}")
    data = json.loads(path.read_text(encoding="utf-8"))
    validate_payload(data)
    return data


def validate_payload(data: dict[str, Any]) -> None:
    for field in ("report_date", "title", "executive_summary", "trend", "categories"):
        if not data.get(field):
            raise ValueError(f"日报数据缺少字段：{field}")
    try:
        datetime.strptime(data["report_date"], "%Y-%m-%d")
    except ValueError as exc:
        raise ValueError("report_date 必须使用 YYYY-MM-DD 格式") from exc
    if not isinstance(data["categories"], list) or not data["categories"]:
        raise ValueError("categories 必须是非空数组")
    total_items = 0
    for category in data["categories"]:
        if not category.get("name"):
            raise ValueError("每个分类必须包含 name")
        if not isinstance(category.get("items"), list):
            raise ValueError(f"分类 {category.get('name')} 的 items 必须是数组")
        total_items += len(category["items"])
        for item in category["items"]:
            for field in REQUIRED_ITEM_FIELDS:
                if not item.get(field):
                    raise ValueError(f"新闻条目缺少字段：{field}")
    if total_items < 1:
        raise ValueError("日报至少需要 1 条新闻")


def set_font(run, font_name: str = "Arial") -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if "：" in text:
        label, rest = text.split("：", 1)
        r1 = p.add_run(label + "：")
        r1.bold = True
        set_font(r1)
        r2 = p.add_run(rest)
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.line_spacing = 1.15
        for run in paragraph.runs:
            run.font.size = Pt(10.5)


def add_source(paragraph, item: dict[str, Any]) -> None:
    run = paragraph.add_run(f"来源：{item['source']}；日期：{item['date']}；链接：{item['url']}")
    set_font(run)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(110, 110, 110)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    h1 = document.styles["Heading 1"]
    h1.font.name = "Arial"
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(31, 78, 121)

    h2 = document.styles["Heading 2"]
    h2.font.name = "Arial"
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(54, 95, 145)


def add_bold_terms(paragraph, text: str, terms: tuple[str, ...]) -> None:
    cursor = text
    while cursor:
        positions = [(cursor.find(term), term) for term in terms if cursor.find(term) >= 0]
        if not positions:
            run = paragraph.add_run(cursor)
            set_font(run)
            return
        idx, term = min(positions, key=lambda pair: pair[0])
        if idx:
            run = paragraph.add_run(cursor[:idx])
            set_font(run)
        run = paragraph.add_run(term)
        set_font(run)
        run.bold = True
        cursor = cursor[idx + len(term):]


def build_docx(data: dict[str, Any], output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    report_date = datetime.strptime(data["report_date"], "%Y-%m-%d").date()
    docx_path = output_dir / f"外贸资讯日报_{report_date:%Y-%m-%d}.docx"

    document = Document()
    configure_document(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)
    run = title.add_run(data["title"])
    set_font(run)
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"生成日期：{report_date:%Y年%m月%d日}")
    set_font(run)
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(90, 90, 90)
    document.add_page_break()

    document.add_heading("执行摘要", level=1)
    p = document.add_paragraph()
    add_bold_terms(p, data["executive_summary"], ("关税", "制裁", "油价", "汇率", "美元", "物流", "锁汇"))

    document.add_heading("重点新闻分类", level=1)
    for category in data["categories"]:
        document.add_heading(category["name"], level=2)
        for index, item in enumerate(category["items"], start=1):
            heading = document.add_paragraph()
            heading.paragraph_format.space_before = Pt(8)
            heading.paragraph_format.space_after = Pt(4)
            run = heading.add_run(f"{index}. {item['title']}")
            set_font(run)
            run.bold = True
            run.font.size = Pt(12)

            table = document.add_table(rows=3, cols=1)
            table.style = "Table Grid"
            set_cell_text(table.cell(0, 0), f"核心内容：{item['summary']}")
            shade_cell(table.cell(0, 0), "F4F8FB")
            set_cell_text(table.cell(1, 0), f"潜在影响：{item['impact']}")
            table.cell(2, 0).text = ""
            add_source(table.cell(2, 0).paragraphs[0], item)

    document.add_heading("市场趋势研判", level=1)
    p = document.add_paragraph()
    run = p.add_run(data["trend"])
    set_font(run)

    document.add_heading("新闻来源清单", level=1)
    for category in data["categories"]:
        for item in category["items"]:
            p = document.add_paragraph()
            run = p.add_run(f"{item['date']}｜{item['source']}｜{item['title']}｜{item['url']}")
            set_font(run)
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(90, 90, 90)

    document.save(docx_path)
    return docx_path


def gmail_credentials(env_file: Path) -> tuple[str, str]:
    load_dotenv(env_file)
    address = os.getenv("GMAIL_ADDRESS", "").strip()
    password = "".join(os.getenv("GMAIL_APP_PASSWORD", "").split())
    missing = [name for name, value in (("GMAIL_ADDRESS", address), ("GMAIL_APP_PASSWORD", password)) if not value]
    if missing:
        raise RuntimeError("缺少 Gmail 配置：" + "、".join(missing))
    return address, password


def send_email(docx_path: Path, data: dict[str, Any], env_file: Path) -> None:
    address, password = gmail_credentials(env_file)
    report_date = datetime.strptime(data["report_date"], "%Y-%m-%d").date()
    subject = data.get("email_subject") or f"【外贸日报】{report_date:%Y年%m月%d日} 全球外贸资讯总结"

    msg = EmailMessage()
    msg["From"] = address
    msg["To"] = address
    msg["Subject"] = subject
    msg.set_content(data["executive_summary"])
    msg.add_attachment(
        docx_path.read_bytes(),
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=docx_path.name,
    )

    context = ssl.create_default_context()
    with smtplib.SMTP_SSL("smtp.gmail.com", 465, context=context) as server:
        server.login(address, password)
        server.send_message(msg)
    print(f"✅ 邮件已发送至：{address}")


def main() -> None:
    args = parse_args()
    try:
        if args.deepseek_research:
            print("步骤1/4：正在调用 DeepSeek API 检索并总结外贸资讯。")
            data = call_deepseek_research(args)
            if args.save_payload:
                args.save_payload.parent.mkdir(parents=True, exist_ok=True)
                args.save_payload.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
                print(f"✅ DeepSeek 日报数据已保存：{args.save_payload}")
        else:
            if not args.payload:
                raise RuntimeError("请提供 payload JSON，或使用 --deepseek-research 调用 DeepSeek API 生成日报数据。")
            print("步骤1/4：正在读取日报数据。")
            data = load_payload(args.payload)

        print("步骤2/4：正在生成 Word 文档。")
        docx_path = build_docx(data, args.output_dir)
        print(f"✅ 文档已保存：{docx_path}")
        print("步骤3/4：Word 文档生成完成。")
        if args.send_email:
            print("步骤4/4：正在通过 Gmail SMTP 发送邮件。")
            send_email(docx_path, data, args.env_file)
        else:
            print("步骤4/4：未请求邮件发送，已跳过。")
    except Exception as exc:
        print(f"❌ 执行失败：{exc}")
        raise


if __name__ == "__main__":
    main()
